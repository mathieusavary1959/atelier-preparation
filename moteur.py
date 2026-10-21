# -*- coding: utf-8 -*-
"""
Moteur documentaire de l'Atelier de préparation.
Extraction de texte (PDF, Word, texte), découpage, indexation locale
et recherche par pertinence, sans aucun service externe.
La génération par IA (clé API) est branchée dans app.py par-dessus
les extraits trouvés ici : les réponses restent ancrées dans les documents.
"""
import io
import json
import math
import os
import re
import unicodedata
from collections import Counter

DOSSIER = os.path.join(os.path.dirname(os.path.abspath(__file__)), "donnees")
FICHIER_INDEX = os.path.join(DOSSIER, "index.json")
DOSSIER_DOCS = os.path.join(DOSSIER, "documents")

MOTS_VIDES = set("""
le la les de des du un une et est sont au aux en dans sur par pour avec sans
qui que quoi dont où ce cet cette ces il elle ils elles nous vous on se sa son
ses leur leurs a à y d l j n s t c m qu aux être avoir fait faire plus moins
très aussi comme mais ou donc car ne pas au aux être auxiliaire the of and to
""".split())


def plier(texte):
    """Minuscules sans accents, pour comparer sans rater les correspondances."""
    texte = unicodedata.normalize("NFD", texte.lower())
    return "".join(c for c in texte if unicodedata.category(c) != "Mn")


def jetons(texte):
    mots = re.findall(r"[a-zàâäéèêëîïôöùûüç0-9]{3,}", plier(texte))
    return [m for m in mots if m not in MOTS_VIDES]


# ---------------------------------------------------------------- extraction
def extraire_pages(nom, octets):
    """Retourne une liste de (page, texte). Page = numéro réel pour un PDF,
    bloc de paragraphes pour un Word."""
    bas = nom.lower()
    if bas.endswith(".pdf"):
        import fitz
        doc = fitz.open(stream=octets, filetype="pdf")
        pages = []
        for i, page in enumerate(doc):
            pages.append((i + 1, page.get_text()))
        doc.close()
        return pages
    if bas.endswith(".docx"):
        import docx
        d = docx.Document(io.BytesIO(octets))
        paragraphes = [p.text for p in d.paragraphs if p.text.strip()]
        for tableau in d.tables:
            for ligne in tableau.rows:
                for cellule in ligne.cells:
                    if cellule.text.strip():
                        paragraphes.append(cellule.text)
        pages, bloc, numero = [], [], 1
        for p in paragraphes:
            bloc.append(p)
            if sum(len(x) for x in bloc) > 2500:
                pages.append((numero, "\n".join(bloc)))
                bloc, numero = [], numero + 1
        if bloc:
            pages.append((numero, "\n".join(bloc)))
        return pages
    if bas.endswith((".txt", ".md")):
        return [(1, octets.decode("utf-8", errors="ignore"))]
    raise ValueError("Format non pris en charge : " + nom)


def decouper(nom_doc, theme, pages, taille=1200, chevauchement=150):
    """Découpe le texte en extraits repérables (document, page)."""
    extraits = []
    for page, texte in pages:
        texte = re.sub(r"\s+", " ", texte).strip()
        debut = 0
        while debut < len(texte):
            fin = debut + taille
            extraits.append({
                "doc": nom_doc,
                "theme": theme,
                "page": page,
                "texte": texte[debut:fin],
            })
            if fin >= len(texte):
                break
            debut = fin - chevauchement
    return extraits


# ---------------------------------------------------------------- index
def charger_index():
    if os.path.exists(FICHIER_INDEX):
        with open(FICHIER_INDEX, encoding="utf-8") as f:
            return json.load(f)
    return {"documents": [], "extraits": []}


def sauver_index(index):
    os.makedirs(DOSSIER_DOCS, exist_ok=True)
    with open(FICHIER_INDEX, "w", encoding="utf-8") as f:
        json.dump(index, f, ensure_ascii=False)


def ajouter_document(index, nom, theme, octets):
    os.makedirs(DOSSIER_DOCS, exist_ok=True)
    chemin = os.path.join(DOSSIER_DOCS, nom)
    with open(chemin, "wb") as f:
        f.write(octets)
    pages = extraire_pages(nom, octets)
    nouveaux = decouper(nom, theme, pages)
    index["extraits"] = [e for e in index["extraits"] if e["doc"] != nom]
    index["extraits"].extend(nouveaux)
    index["documents"] = [d for d in index["documents"] if d["nom"] != nom]
    index["documents"].append({
        "nom": nom, "theme": theme,
        "pages": len(pages), "extraits": len(nouveaux),
    })
    sauver_index(index)
    return len(pages), len(nouveaux)


def retirer_document(index, nom):
    index["extraits"] = [e for e in index["extraits"] if e["doc"] != nom]
    index["documents"] = [d for d in index["documents"] if d["nom"] != nom]
    chemin = os.path.join(DOSSIER_DOCS, nom)
    if os.path.exists(chemin):
        os.remove(chemin)
    sauver_index(index)


# ---------------------------------------------------------------- recherche
def rechercher(index, question, k=5, theme=None):
    """Score de pertinence maison (fréquence des termes, rareté dans le fonds)."""
    q = Counter(jetons(question))
    if not q or not index["extraits"]:
        return []
    fonds = [e for e in index["extraits"] if theme in (None, "Tous", e["theme"])]
    n = max(len(fonds), 1)
    df = Counter()
    for e in fonds:
        for m in set(jetons(e["texte"])):
            df[m] += 1
    resultats = []
    for e in fonds:
        compteur = Counter(jetons(e["texte"]))
        score = 0.0
        for mot, fois in q.items():
            if compteur[mot]:
                idf = math.log((n + 1) / (df[mot] + 1)) + 1
                score += fois * (1 + math.log(compteur[mot])) * idf
        if score > 0:
            resultats.append((score, e))
    resultats.sort(key=lambda r: -r[0])
    return resultats[:k]


def phrases_pertinentes(texte, question, max_phrases=3):
    q = set(jetons(question))
    phrases = re.split(r"(?<=[.!?])\s+", texte)
    notees = []
    for p in phrases:
        p = p.strip()
        if len(p) < 25:
            continue
        hits = sum(1 for m in jetons(p) if m in q)
        notees.append((hits, p))
    notees.sort(key=lambda r: -r[0])
    return [p for h, p in notees if h > 0][:max_phrases]


def reponse_documentaire(question, resultats):
    """Réponse assemblée uniquement avec des phrases des documents,
    chacune suivie de sa source. Aucune invention."""
    if not resultats:
        return ("Je ne trouve rien dans tes documents sur ce sujet. "
                "Vérifie le thème choisi ou dépose le document pertinent.")
    lignes = ["D'après tes documents :"]
    sources = []
    for score, e in resultats:
        phrases = phrases_pertinentes(e["texte"], question)
        if not phrases and len(sources) == 0:
            phrases = [e["texte"][:300].strip()]
        for p in phrases[:2]:
            lignes.append("- " + p + f"  [{e['doc']}, p. {e['page']}]")
        s = f"{e['doc']}, p. {e['page']}"
        if s not in sources:
            sources.append(s)
    lignes.append("")
    lignes.append("Sources : " + " ; ".join(sources))
    return "\n".join(lignes)
